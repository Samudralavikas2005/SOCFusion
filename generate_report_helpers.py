import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_header_styled(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    if level == 1:
        run = h.runs[0]
        run.font.color.rgb = RGBColor(16, 44, 87) # Deep Navy
        run.font.size = Pt(20)
        run.font.bold = True
    elif level == 2:
        run = h.runs[0]
        run.font.color.rgb = RGBColor(41, 75, 138) # Slate Blue
        run.font.size = Pt(15)
        run.font.bold = True
    elif level == 3:
        run = h.runs[0]
        run.font.color.rgb = RGBColor(60, 64, 72) # Charcoal
        run.font.size = Pt(12)
        run.font.bold = True
    return h

def add_p(doc, text, bold_prefix="", space_after=6, line_spacing=1.15):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.size = Pt(11)
        r_pre.font.name = "Calibri"
    r = p.add_run(text)
    r.font.size = Pt(11)
    r.font.name = "Calibri"
    return p

def add_callout(doc, text, title=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border accent
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="102C57"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        r_title = p.add_run(f"NOTE / KEY TAKEAWAY: {title}\n")
        r_title.bold = True
        r_title.font.color.rgb = RGBColor(16, 44, 87)
        r_title.font.size = Pt(10.5)
    r_body = p.add_run(text)
    r_body.font.size = Pt(10.5)
    r_body.font.name = "Calibri"
    r_body.font.italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "1E1E1E") # Dark IDE background
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(code_text)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(220, 220, 220)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def format_table(table, col_widths, headers, data=None, header_bg="102C57"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.bold = True
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.name = "Calibri"
            
    # Data rows
    if data:
        for row_idx, row_data in enumerate(data):
            row = table.add_row()
            cells = row.cells
            bg_color = "F9FBFD" if row_idx % 2 == 1 else "FFFFFF"
            for c_idx, val in enumerate(row_data):
                cells[c_idx].text = str(val)
                set_cell_background(cells[c_idx], bg_color)
                set_cell_margins(cells[c_idx], top=80, bottom=80, left=120, right=120)
                p = cells[c_idx].paragraphs[0]
                for r in p.runs:
                    r.font.size = Pt(9.5)
                    r.font.name = "Calibri"
                
    # Column widths
    for row in table.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)

print("Report helper modules initialized.")

